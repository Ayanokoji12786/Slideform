import logging
import math
import os
import re
import shutil
import subprocess
import tempfile
import threading
import traceback
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill


# ============================================================
# SETTINGS
# ============================================================

APP_TITLE = "PowerPoint to Excel Converter"

# How wide the slide appears inside Excel.
DISPLAY_WIDTH = 960

LOG_FILE = "ppt_to_excel.log"
DEFAULT_SLIDES_TO_EXTRACT = ""
DEFAULT_TABLE_SEARCH_TERM = "Demand Priority"

PPTX_NAMESPACES = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}


# ============================================================
# LOGGING
# ============================================================

logging.basicConfig(
    filename=LOG_FILE,
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(message)s",
)


# ============================================================
# POWERPOINT EXPORT
# ============================================================

def export_slides_to_png(
    pptx_path,
    temp_dir,
    progress_callback=None
):
    """Use PowerPoint for Mac's AppleScript interface to export all slides."""

    if shutil.which("osascript") is None:
        raise RuntimeError("macOS osascript was not found.")

    powerpoint_app = Path("/Applications/Microsoft PowerPoint.app")
    if not powerpoint_app.exists():
        raise RuntimeError(
            "Microsoft PowerPoint for Mac was not found in Applications."
        )

    # Exporting each slide separately is more dependable than PowerPoint for
    # Mac's bulk ``save presentation as PNG`` command. Passing paths as argv
    # rather than interpolating them into AppleScript safely handles spaces,
    # quotes, and non-English filenames.
    script = '''on run argv
    set inputFile to POSIX file (item 1 of argv)
    set exportFolder to item 2 of argv
    tell application "Microsoft PowerPoint"
        activate
        open inputFile
        set thePresentation to active presentation
        try
            set slideCount to count of slides of thePresentation
            if slideCount is less than 1 then error "The PowerPoint presentation has no slides."
            repeat with slideIndex from 1 to slideCount
                set outputFile to exportFolder & "/slide_" & slideIndex & ".png"
                set theSlide to slide slideIndex of thePresentation
                save theSlide in outputFile as save as PNG
            end repeat
        on error errorMessage number errorNumber
            try
                close thePresentation saving no
            end try
            error errorMessage number errorNumber
        end try
        close thePresentation saving no
    end tell
    return slideCount
end run'''

    if progress_callback:
        progress_callback("Opening PowerPoint and exporting slides...", 0, 1)

    logging.info("Exporting PowerPoint file with AppleScript: %s", pptx_path)
    try:
        result = subprocess.run(
            ["osascript", "-e", script, str(Path(pptx_path).resolve()), temp_dir],
            capture_output=True,
            text=True,
            timeout=300,
            check=False,
        )
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError(
            "PowerPoint did not finish exporting within five minutes."
        ) from exc

    if result.returncode != 0:
        details = (result.stderr or result.stdout).strip()
        logging.error("PowerPoint AppleScript failed: %s", details)
        raise RuntimeError(
            "PowerPoint could not export the slides. "
            "Check macOS Automation permission, then try again.\n\n"
            f"Details: {details or 'Unknown AppleScript error.'}"
        )

    try:
        expected_slide_count = int(result.stdout.strip())
    except ValueError as exc:
        raise RuntimeError(
            "PowerPoint completed without reporting how many slides it exported."
        ) from exc

    def slide_sort_key(path):
        """Keep PowerPoint's Slide1, Slide2, ... ordering (not Slide1, Slide10)."""
        numbers = re.findall(r"\d+", path.stem)
        return (int(numbers[-1]) if numbers else 0, path.name.casefold())

    exported_files = sorted(
        (
            path for path in Path(temp_dir).rglob("*")
            if path.is_file() and path.suffix.casefold() == ".png"
        ),
        key=slide_sort_key,
    )
    if not exported_files:
        raise RuntimeError(
            "PowerPoint did not create any PNG slide images. The presentation "
            "may be invalid or PowerPoint's PNG export failed."
        )

    if len(exported_files) != expected_slide_count:
        raise RuntimeError(
            f"PowerPoint exported {len(exported_files)} of "
            f"{expected_slide_count} slides."
        )

    if progress_callback:
        progress_callback("PowerPoint slide export completed.", 1, 1)

    logging.info("Successfully exported %d slides.", len(exported_files))
    return [str(path) for path in exported_files]


# ============================================================
# EXCEL IMAGE INSERTION
# ============================================================

def add_slide_image(
    worksheet,
    image_path,
    anchor
):
    """
    Adds a slide PNG to Excel while preserving aspect ratio.
    """

    image = ExcelImage(image_path)

    original_width = float(image.width)
    original_height = float(image.height)

    if original_width <= 0 or original_height <= 0:
        raise RuntimeError(
            f"Invalid image dimensions: {image_path}"
        )

    # Resize for display in Excel,
    # but preserve the aspect ratio.
    image.width = DISPLAY_WIDTH

    image.height = int(
        round(
            DISPLAY_WIDTH *
            original_height /
            original_width
        )
    )

    worksheet.add_image(
        image,
        anchor
    )

    return image.width, image.height


# ============================================================
# EXCEL CREATION
# ============================================================

def parse_slide_numbers(value):
    """Parse a comma-separated slide list without guessing any slide numbers."""
    values = [item.strip() for item in value.split(",") if item.strip()]
    if not values:
        raise ValueError("Enter at least one slide number, for example: 6 or 12, 13, 15.")

    slide_numbers = []
    for item in values:
        if not item.isdigit() or int(item) < 1:
            raise ValueError(
                "Slide numbers must be positive whole numbers separated by commas, "
                "for example: 12, 13, 15."
            )
        number = int(item)
        if number in slide_numbers:
            raise ValueError(f"Slide {number} was entered more than once.")
        slide_numbers.append(number)

    return slide_numbers


def get_powerpoint_table_text(cell):
    """Return the native PowerPoint table cell text, including intentional blanks."""
    paragraphs = []
    for paragraph in cell.findall(".//a:p", PPTX_NAMESPACES):
        text_runs = [
            run.text or ""
            for run in paragraph.findall(".//a:t", PPTX_NAMESPACES)
        ]
        paragraphs.append("".join(text_runs))
    return "\n".join(paragraphs)


def normalize_powerpoint_text(value):
    """Compare native PowerPoint text without changing the values copied to Excel."""
    return re.sub(r"\s+", " ", value).strip().casefold()


def get_text_body_text(text_body):
    """Read one PowerPoint text box as displayed, preserving paragraph breaks."""
    paragraphs = []
    for paragraph in text_body.findall("a:p", PPTX_NAMESPACES):
        paragraphs.append(
            "".join(
                run.text or ""
                for run in paragraph.findall(".//a:t", PPTX_NAMESPACES)
            )
        )
    return "\n".join(paragraphs)


def parse_slide_xml(data):
    """Parse one slide's XML, rejecting a DOCTYPE, which real PowerPoint slide parts never
    contain. Guards against entity-expansion ("billion laughs") denial of service from a
    crafted .pptx, since ElementTree does not disable that by default."""
    if b"<!DOCTYPE" in data:
        raise ElementTree.ParseError("DOCTYPE declarations are not allowed in PowerPoint slide XML.")
    return ElementTree.fromstring(data)


def find_slides_containing_text(pptx_path, search_term=DEFAULT_TABLE_SEARCH_TERM):
    """Find slide numbers whose native text or table rows contain a requested phrase."""
    needle = normalize_powerpoint_text(search_term)
    if not needle:
        raise ValueError("The default slide search text cannot be empty.")

    matching_slides = []
    try:
        with zipfile.ZipFile(pptx_path) as presentation:
            slide_files = []
            for name in presentation.namelist():
                match = re.fullmatch(r"ppt/slides/slide(\d+)\.xml", name)
                if match:
                    slide_files.append((int(match.group(1)), name))

            for slide_number, slide_file in sorted(slide_files):
                slide_xml = parse_slide_xml(presentation.read(slide_file))

                text_boxes = [
                    get_text_body_text(text_body)
                    for text_body in slide_xml.findall(".//a:txBody", PPTX_NAMESPACES)
                ]
                table_rows = [
                    " ".join(
                        get_powerpoint_table_text(cell)
                        for cell in row.findall("a:tc", PPTX_NAMESPACES)
                    )
                    for row in slide_xml.findall(".//a:tbl/a:tr", PPTX_NAMESPACES)
                ]

                if any(
                    needle in normalize_powerpoint_text(value)
                    for value in text_boxes + table_rows
                ):
                    matching_slides.append(slide_number)
    except zipfile.BadZipFile as exc:
        raise ValueError("The selected file is not a valid .pptx PowerPoint file.") from exc
    except ElementTree.ParseError as exc:
        raise ValueError("PowerPoint slide text could not be read.") from exc

    return matching_slides


def extract_powerpoint_tables(pptx_path, slide_numbers, extract_all_tables):
    """Copy native PowerPoint table cells; never derive values from slide graphics."""
    extracted_tables = []

    try:
        with zipfile.ZipFile(pptx_path) as presentation:
            slide_files = set(presentation.namelist())

            for slide_number in slide_numbers:
                slide_file = f"ppt/slides/slide{slide_number}.xml"
                if slide_file not in slide_files:
                    raise ValueError(
                        f"Slide {slide_number} does not exist in this PowerPoint file."
                    )

                slide_xml = parse_slide_xml(presentation.read(slide_file))
                tables = slide_xml.findall(".//a:tbl", PPTX_NAMESPACES)
                if not tables:
                    raise ValueError(
                        f"Slide {slide_number} does not contain a native PowerPoint table."
                    )

                selected_tables = tables if extract_all_tables else tables[:1]
                for table_number, table in enumerate(selected_tables, start=1):
                    rows = []
                    for row in table.findall("a:tr", PPTX_NAMESPACES):
                        rows.append(
                            [
                                get_powerpoint_table_text(cell)
                                for cell in row.findall("a:tc", PPTX_NAMESPACES)
                            ]
                        )
                    if rows:
                        extracted_tables.append((slide_number, table_number, rows))

    except zipfile.BadZipFile as exc:
        raise ValueError("The selected file is not a valid .pptx PowerPoint file.") from exc
    except ElementTree.ParseError as exc:
        raise ValueError("PowerPoint contains a table that could not be read.") from exc

    if not extracted_tables:
        raise ValueError("No PowerPoint table data was found for the selected slides.")

    return extracted_tables


def create_table_workbook(table_data, output_path, progress_callback=None):
    """Create one worksheet per selected native PowerPoint table."""
    workbook = Workbook()

    try:
        for index, (slide_number, table_number, rows) in enumerate(table_data, start=1):
            if progress_callback:
                progress_callback(
                    f"Copying table {index} of {len(table_data)} to Excel...",
                    index - 1,
                    len(table_data),
                )

            worksheet = workbook.active if index == 1 else workbook.create_sheet()
            worksheet.title = f"Slide {slide_number} Table {table_number}"
            worksheet.sheet_view.showGridLines = False
            worksheet.freeze_panes = "A2"

            for row in rows:
                worksheet.append(row)

            # The first PowerPoint table row becomes the Excel header row.
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(wrap_text=True, vertical="top")

            for column_cells in worksheet.iter_cols():
                width = max(
                    (len(str(cell.value or "").split("\n")[0]) for cell in column_cells),
                    default=0,
                )
                worksheet.column_dimensions[column_cells[0].column_letter].width = min(
                    max(width + 2, 12), 50
                )
                for cell in column_cells:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")

        if progress_callback:
            progress_callback("Saving Excel workbook...", len(table_data), len(table_data))

        workbook.save(os.path.abspath(output_path))
        logging.info("Saved extracted table workbook: %s", output_path)
    finally:
        workbook.close()

def create_excel_workbook(
    image_paths,
    output_path,
    mode,
    progress_callback=None
):
    """
    Creates the Excel workbook.
    """

    workbook = Workbook()

    try:

        slide_count = len(image_paths)

        if slide_count == 0:
            raise RuntimeError(
                "No slide images were created."
            )

        # ====================================================
        # MODE 1:
        # Each slide gets its own Excel worksheet.
        # ====================================================

        if mode == "Separate worksheets":

            first_sheet = workbook.active
            first_sheet.title = "Slide 1"

            for index, image_path in enumerate(
                image_paths,
                start=1
            ):

                if progress_callback:
                    progress_callback(
                        f"Adding slide "
                        f"{index} of {slide_count} "
                        f"to Excel...",
                        index - 1,
                        slide_count,
                    )

                if index == 1:
                    worksheet = first_sheet
                else:
                    worksheet = workbook.create_sheet(
                        title=f"Slide {index}"
                    )

                worksheet.sheet_view.showGridLines = False

                # Make column A wide.
                worksheet.column_dimensions["A"].width = 135

                # Add slide starting at A1.
                add_slide_image(
                    worksheet,
                    image_path,
                    "A1"
                )

        # ====================================================
        # MODE 2:
        # All slides vertically on one worksheet.
        # ====================================================

        elif mode == "All slides vertically":

            worksheet = workbook.active
            worksheet.title = "Slides"

            worksheet.sheet_view.showGridLines = False
            worksheet.column_dimensions["A"].width = 135

            current_row = 1

            for index, image_path in enumerate(
                image_paths,
                start=1
            ):

                if progress_callback:
                    progress_callback(
                        f"Adding slide "
                        f"{index} of {slide_count} "
                        f"to Excel...",
                        index - 1,
                        slide_count,
                    )

                _, image_height = add_slide_image(
                    worksheet,
                    image_path,
                    f"A{current_row}"
                )

                # Roughly 20 pixels per default Excel row.
                rows_used = max(
                    1,
                    math.ceil(image_height / 20)
                )

                # Leave space between slides.
                current_row += rows_used + 3

        else:

            raise ValueError(
                f"Unknown conversion mode: {mode}"
            )

        # ====================================================
        # SAVE
        # ====================================================

        if progress_callback:
            progress_callback(
                "Saving Excel workbook...",
                slide_count,
                slide_count
            )

        output_path = os.path.abspath(output_path)

        workbook.save(output_path)

        logging.info(
            "Saved Excel workbook: %s",
            output_path
        )

    finally:

        try:
            workbook.close()
        except Exception:
            pass


# ============================================================
# COMPLETE CONVERSION PIPELINE
# ============================================================

def convert_powerpoint_to_excel(
    pptx_path,
    output_path,
    conversion_type,
    image_mode,
    slide_numbers_text,
    extract_all_tables,
    progress_callback=None
):
    """
    Convert native table data or slide images into Excel.
    """

    try:

        # ----------------------------------------------------
        # Validate PowerPoint input.
        # ----------------------------------------------------

        if not pptx_path:
            raise ValueError(
                "Please select a PowerPoint file."
            )

        pptx = Path(pptx_path)

        if not pptx.exists():
            raise FileNotFoundError(
                "The selected PowerPoint file does not exist."
            )

        if pptx.suffix.lower() != ".pptx":
            raise ValueError(
                "Please select a .pptx PowerPoint file."
            )

        # ----------------------------------------------------
        # Validate Excel output.
        # ----------------------------------------------------

        if not output_path:
            raise ValueError(
                "Please select where to save the Excel file."
            )

        output = Path(output_path)

        if output.suffix.lower() != ".xlsx":
            raise ValueError(
                "The Excel output file must end with .xlsx."
            )

        output.parent.mkdir(
            parents=True,
            exist_ok=True
        )

        if conversion_type == "Extract table data":
            slide_numbers = parse_slide_numbers(slide_numbers_text)
            if progress_callback:
                progress_callback("Reading selected PowerPoint tables...", 0, 1)

            table_data = extract_powerpoint_tables(
                str(pptx), slide_numbers, extract_all_tables
            )
            create_table_workbook(table_data, str(output), progress_callback)

        elif conversion_type == "Insert slide images":
            if progress_callback:
                progress_callback("Opening PowerPoint...", 0, 1)

            # Temporary directory automatically disappears after conversion.
            with tempfile.TemporaryDirectory(prefix="ppt_to_excel_") as temp_dir:
                image_paths = export_slides_to_png(
                    str(pptx), temp_dir, progress_callback
                )
                create_excel_workbook(
                    image_paths, str(output), image_mode, progress_callback
                )

        else:
            raise ValueError(f"Unknown conversion type: {conversion_type}")

        if progress_callback:
            progress_callback(
                "Conversion completed successfully.",
                1,
                1
            )

    except OSError as exc:
        raise RuntimeError(f"Could not run PowerPoint automation: {exc}") from exc


# ============================================================
# GUI
# ============================================================

class ConverterApp:

    def __init__(self, root):

        self.root = root

        self.root.title(APP_TITLE)
        self.root.geometry("700x510")
        self.root.minsize(640, 480)

        self.pptx_path = tk.StringVar()

        self.output_path = tk.StringVar()

        self.conversion_type = tk.StringVar(
            value="Extract table data"
        )

        self.slide_numbers = tk.StringVar(
            value=DEFAULT_SLIDES_TO_EXTRACT
        )

        self.table_selection = tk.StringVar(
            value="First table on each selected slide"
        )

        self.image_mode = tk.StringVar(
            value="Separate worksheets"
        )

        self.status = tk.StringVar(
            value="Choose a PowerPoint file to begin."
        )

        self.build_ui()


    def build_ui(self):

        outer = ttk.Frame(
            self.root,
            padding=20
        )

        outer.pack(
            fill="both",
            expand=True
        )

        # ----------------------------------------------------
        # TITLE
        # ----------------------------------------------------

        title = ttk.Label(
            outer,
            text=APP_TITLE,
            font=(
                "Segoe UI",
                18,
                "bold"
            ),
        )

        title.pack(
            anchor="w",
            pady=(0, 20)
        )

        # ----------------------------------------------------
        # POWERPOINT FILE
        # ----------------------------------------------------

        ppt_frame = ttk.Frame(outer)

        ppt_frame.pack(
            fill="x",
            pady=6
        )

        ttk.Label(
            ppt_frame,
            text="PowerPoint file:"
        ).pack(
            anchor="w"
        )

        ppt_row = ttk.Frame(ppt_frame)

        ppt_row.pack(
            fill="x",
            pady=(5, 0)
        )

        self.ppt_entry = ttk.Entry(
            ppt_row,
            textvariable=self.pptx_path,
            state="readonly"
        )

        self.ppt_entry.pack(
            side="left",
            fill="x",
            expand=True
        )

        self.ppt_button = ttk.Button(
            ppt_row,
            text="Choose PowerPoint",
            command=self.choose_powerpoint
        )

        self.ppt_button.pack(
            side="left",
            padx=(8, 0)
        )

        # ----------------------------------------------------
        # EXCEL OUTPUT
        # ----------------------------------------------------

        output_frame = ttk.Frame(outer)

        output_frame.pack(
            fill="x",
            pady=6
        )

        ttk.Label(
            output_frame,
            text="Excel output:"
        ).pack(
            anchor="w"
        )

        output_row = ttk.Frame(output_frame)

        output_row.pack(
            fill="x",
            pady=(5, 0)
        )

        self.output_entry = ttk.Entry(
            output_row,
            textvariable=self.output_path,
            state="readonly"
        )

        self.output_entry.pack(
            side="left",
            fill="x",
            expand=True
        )

        self.output_button = ttk.Button(
            output_row,
            text="Choose Output",
            command=self.choose_output
        )

        self.output_button.pack(
            side="left",
            padx=(8, 0)
        )

        # ----------------------------------------------------
        # CONVERSION OPTIONS
        # ----------------------------------------------------

        options_frame = ttk.Frame(outer)

        options_frame.pack(
            fill="x",
            pady=(14, 6)
        )

        ttk.Label(
            options_frame,
            text="Convert:"
        ).pack(
            anchor="w"
        )

        self.conversion_type_box = ttk.Combobox(
            options_frame,
            textvariable=self.conversion_type,
            values=(
                "Extract table data",
                "Insert slide images",
            ),
            state="readonly",
            width=28,
        )

        self.conversion_type_box.pack(
            anchor="w",
            pady=(5, 0)
        )

        ttk.Label(
            options_frame,
            text="Slides to extract (comma-separated):"
        ).pack(anchor="w", pady=(12, 0))

        self.slide_numbers_entry = ttk.Entry(
            options_frame,
            textvariable=self.slide_numbers,
            width=28,
        )
        self.slide_numbers_entry.pack(anchor="w", pady=(5, 0))

        ttk.Label(
            options_frame,
            text='Automatic default: slides containing "Demand Priority".'
        ).pack(anchor="w", pady=(3, 0))

        ttk.Label(options_frame, text="Tables to copy:").pack(
            anchor="w", pady=(12, 0)
        )
        self.table_selection_box = ttk.Combobox(
            options_frame,
            textvariable=self.table_selection,
            values=(
                "First table on each selected slide",
                "All native tables on each selected slide",
            ),
            state="readonly",
            width=36,
        )
        self.table_selection_box.pack(anchor="w", pady=(5, 0))

        ttk.Label(options_frame, text="Slide image layout (when selected):").pack(
            anchor="w", pady=(12, 0)
        )
        self.image_mode_box = ttk.Combobox(
            options_frame,
            textvariable=self.image_mode,
            values=("Separate worksheets", "All slides vertically"),
            state="readonly",
            width=28,
        )
        self.image_mode_box.pack(anchor="w", pady=(5, 0))

        # ----------------------------------------------------
        # PROGRESS
        # ----------------------------------------------------

        self.progress = ttk.Progressbar(
            outer,
            mode="determinate",
            maximum=100
        )

        self.progress.pack(
            fill="x",
            pady=(20, 8)
        )

        ttk.Label(
            outer,
            textvariable=self.status,
            wraplength=640,
        ).pack(
            anchor="w"
        )

        # ----------------------------------------------------
        # CONVERT BUTTON
        # ----------------------------------------------------

        self.convert_button = ttk.Button(
            outer,
            text="Convert",
            command=self.start_conversion
        )

        self.convert_button.pack(
            anchor="e",
            pady=(18, 0)
        )


    # ========================================================
    # CHOOSE POWERPOINT
    # ========================================================

    def choose_powerpoint(self):

        path = filedialog.askopenfilename(
            title="Choose a PowerPoint presentation",
            filetypes=[
                (
                    "PowerPoint Presentation",
                    "*.pptx"
                )
            ],
        )

        if not path:
            return

        self.pptx_path.set(path)

        try:
            matching_slides = find_slides_containing_text(path)
        except ValueError as exc:
            logging.warning("Could not find Demand Priority slides: %s", exc)
            self.slide_numbers.set("")
            self.status.set("PowerPoint selected. Enter slide numbers to extract.")
        else:
            self.slide_numbers.set(
                ", ".join(str(number) for number in matching_slides)
            )
            if matching_slides:
                self.status.set(
                    'Ready to extract slides containing "Demand Priority".'
                )
            else:
                self.status.set(
                    'No slides containing "Demand Priority" were found. '
                    "Enter slide numbers manually."
                )

        # Preserve the PowerPoint filename; only change .pptx to .xlsx.
        suggested = str(Path(path).with_suffix(".xlsx"))

        self.output_path.set(suggested)

    # ========================================================
    # CHOOSE EXCEL OUTPUT
    # ========================================================

    def choose_output(self):

        initial_name = "presentation.xlsx"

        if self.pptx_path.get():

            initial_name = Path(self.pptx_path.get()).with_suffix(".xlsx").name

        path = filedialog.asksaveasfilename(
            title="Save Excel workbook",
            defaultextension=".xlsx",
            initialfile=initial_name,
            filetypes=[
                (
                    "Excel Workbook",
                    "*.xlsx"
                )
            ],
        )

        if path:
            self.output_path.set(path)


    # ========================================================
    # ENABLE/DISABLE BUTTONS
    # ========================================================

    def set_controls_enabled(
        self,
        enabled
    ):

        button_state = (
            "normal"
            if enabled
            else "disabled"
        )

        combobox_state = (
            "readonly"
            if enabled
            else "disabled"
        )

        self.ppt_button.config(
            state=button_state
        )

        self.output_button.config(
            state=button_state
        )

        self.convert_button.config(
            state=button_state
        )

        self.conversion_type_box.config(
            state=combobox_state
        )

        self.table_selection_box.config(state=combobox_state)
        self.image_mode_box.config(state=combobox_state)
        self.slide_numbers_entry.config(state=button_state)


    # ========================================================
    # PROGRESS CALLBACK
    # ========================================================

    def update_progress(
        self,
        text,
        current,
        total
    ):

        def update():

            self.status.set(text)

            if total > 0:

                percent = int(
                    (current / total) * 100
                )

                percent = max(
                    0,
                    min(
                        100,
                        percent
                    )
                )

                self.progress["value"] = percent

            else:

                self.progress["value"] = 0

        self.root.after(
            0,
            update
        )


    # ========================================================
    # START CONVERSION
    # ========================================================

    def start_conversion(self):

        pptx_path = (
            self.pptx_path.get().strip()
        )

        output_path = (
            self.output_path.get().strip()
        )

        conversion_type = self.conversion_type.get()
        image_mode = self.image_mode.get()
        slide_numbers = self.slide_numbers.get().strip()
        extract_all_tables = (
            self.table_selection.get() == "All native tables on each selected slide"
        )

        if not pptx_path:

            messagebox.showwarning(
                APP_TITLE,
                "Please choose a PowerPoint file first."
            )

            return

        if not output_path:

            messagebox.showwarning(
                APP_TITLE,
                "Please choose where to save the Excel file."
            )

            return

        if conversion_type == "Extract table data" and not slide_numbers:
            messagebox.showwarning(
                APP_TITLE,
                "Enter slide numbers to extract, for example: 6 or 12, 13, 15."
            )
            return

        self.set_controls_enabled(False)

        self.progress["value"] = 0

        self.status.set(
            "Starting conversion..."
        )

        # Run conversion separately so the window
        # does not freeze.
        worker = threading.Thread(
            target=self.conversion_worker,
            args=(
                pptx_path,
                output_path,
                conversion_type,
                image_mode,
                slide_numbers,
                extract_all_tables,
            ),
            daemon=True,
        )

        worker.start()


    # ========================================================
    # BACKGROUND WORKER
    # ========================================================

    def conversion_worker(
        self,
        pptx_path,
        output_path,
        conversion_type,
        image_mode,
        slide_numbers,
        extract_all_tables,
    ):

        try:

            convert_powerpoint_to_excel(
                pptx_path,
                output_path,
                conversion_type,
                image_mode,
                slide_numbers,
                extract_all_tables,
                self.update_progress
            )

            self.root.after(
                0,
                lambda:
                    self.conversion_succeeded(
                        output_path
                    )
            )

        except PermissionError:

            logging.exception(
                "Permission error during conversion."
            )

            self.root.after(
                0,
                lambda:
                    self.conversion_failed(
                        "Excel could not save the output file.\n\n"
                        "If the Excel file is already open, "
                        "close it and try again."
                    )
            )

        except Exception as exc:

            logging.error(
                "Conversion failed: %s\n%s",
                exc,
                traceback.format_exc()
            )

            message = (
                self.friendly_error_message(
                    exc
                )
            )

            self.root.after(
                0,
                lambda m=message:
                    self.conversion_failed(m)
            )


    # ========================================================
    # FRIENDLY ERRORS
    # ========================================================

    def friendly_error_message(
        self,
        exc
    ):

        text = str(exc).strip()

        lower = text.lower()

        if (
            "class not registered" in lower
            or
            "invalid class string" in lower
        ):

            return (
                "Microsoft PowerPoint could not be started.\n\n"
                "Make sure the desktop version of "
                "Microsoft PowerPoint is installed."
            )

        if (
            "permission" in lower
            or
            "access is denied" in lower
        ):

            return (
                "macOS denied access to a file or to PowerPoint.\n\n"
                "Close the output workbook if it is open. Then check System "
                "Settings → Privacy & Security → Automation and allow your "
                "Terminal or Python app to control Microsoft PowerPoint."
            )

        if "automation" in lower or "not authorized" in lower:
            return (
                "macOS did not allow this app to control Microsoft PowerPoint.\n\n"
                "Open System Settings → Privacy & Security → Automation, then "
                "enable Microsoft PowerPoint for the app that launched this "
                "program (usually Terminal, iTerm, or Python)."
            )

        if "powerpoint" in lower and "not found" in lower:
            return (
                "Microsoft PowerPoint for Mac could not be found.\n\n"
                "Install PowerPoint in the Applications folder, then try again."
            )

        return (
            "The conversion failed.\n\n"
            +
            (
                text
                if text
                else
                "An unexpected error occurred."
            )
            +
            "\n\nTechnical details were written to "
            +
            LOG_FILE
        )


    # ========================================================
    # SUCCESS
    # ========================================================

    def conversion_succeeded(
        self,
        output_path
    ):

        self.set_controls_enabled(True)

        self.progress["value"] = 100

        self.status.set(
            "Conversion completed successfully."
        )

        messagebox.showinfo(
            APP_TITLE,
            "Done!\n\n"
            "Excel file created at:\n"
            +
            output_path
        )


    # ========================================================
    # FAILURE
    # ========================================================

    def conversion_failed(
        self,
        message
    ):

        self.set_controls_enabled(True)

        self.status.set(
            "Conversion failed."
        )

        messagebox.showerror(
            APP_TITLE,
            message
        )


# ============================================================
# START APPLICATION
# ============================================================

def main():

    logging.info(
        "Application started."
    )

    root = tk.Tk()

    ConverterApp(root)

    root.mainloop()


if __name__ == "__main__":
    main()
